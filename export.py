"""월간 보고서 마크다운(report.py 결과)을 PDF·워드 파일로 바꾼다.

report.py가 만드는 마크다운만 다룬다: # 제목, ## 소제목, - 목록, | 표 |, **굵게**, ※ 주석.
"""
import io
import re
from pathlib import Path

FONT_DIR = Path(__file__).parent / "fonts"
BOLD = re.compile(r"\*\*(.+?)\*\*")


def parse(md: str) -> list:
    """마크다운을 (종류, 내용) 블록 목록으로 나눈다."""
    blocks, table = [], None
    for line in md.splitlines() + [""]:
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if table is None:
                table = {"rows": [], "align": None}
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                table["align"] = ["right" if c.endswith(":") else "left" for c in cells]
            else:
                table["rows"].append(cells)
            continue
        if table is not None:
            blocks.append(("table", table))
            table = None
        if not s:
            continue
        if s.startswith("## "):
            blocks.append(("h2", s[3:]))
        elif s.startswith("# "):
            blocks.append(("h1", s[2:]))
        elif s.startswith("- "):
            blocks.append(("bullet", s[2:]))
        elif s.startswith("※"):
            blocks.append(("note", s))
        else:
            blocks.append(("para", s))
    return blocks


# ---------------------------------------------------------------- PDF (reportlab, 나눔고딕 내장)
def report_pdf(md: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.fonts import addMapping

    if "Nanum" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Nanum", str(FONT_DIR / "NanumGothic.ttf")))
        pdfmetrics.registerFont(TTFont("Nanum-Bold", str(FONT_DIR / "NanumGothicBold.ttf")))
        addMapping("Nanum", 0, 0, "Nanum")
        addMapping("Nanum", 1, 0, "Nanum-Bold")
        addMapping("Nanum", 0, 1, "Nanum")
        addMapping("Nanum", 1, 1, "Nanum-Bold")

    ink, muted, line, head_bg = (colors.HexColor(c) for c in ("#15232d", "#5b6b78", "#d5dde3", "#eef3f5"))
    base = ParagraphStyle("base", fontName="Nanum", fontSize=9.5, leading=14, textColor=ink, wordWrap="CJK")
    st = {
        "h1": ParagraphStyle("h1", base, fontName="Nanum-Bold", fontSize=17, leading=24, spaceAfter=8),
        "h2": ParagraphStyle("h2", base, fontName="Nanum-Bold", fontSize=12.5, leading=18, spaceBefore=12, spaceAfter=6,
                             textColor=colors.HexColor("#157f8d")),
        "para": ParagraphStyle("para", base, spaceBefore=6, spaceAfter=3),
        "bullet": ParagraphStyle("bullet", base, leftIndent=10, bulletIndent=0, spaceAfter=2),
        "note": ParagraphStyle("note", base, fontSize=8.5, textColor=muted, spaceBefore=4),
        "cell": ParagraphStyle("cell", base, fontSize=8.5, leading=11.5),
    }
    st["cell_r"] = ParagraphStyle("cell_r", st["cell"], alignment=TA_RIGHT)

    def rich(text):
        text = text.replace("⚠ ", "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return BOLD.sub(r"<b>\1</b>", text)

    story = []
    width = A4[0] - 36 * mm
    for kind, body in parse(md):
        if kind == "table":
            rows, align = body["rows"], body["align"] or ["left"] * len(body["rows"][0])
            data = [[Paragraph(f'<font color="#d03a3a">{rich(c)}</font>' if "▼" in c or "⚠" in c else rich(c), st["cell_r"] if a == "right" and i else st["cell"])
                     for c, a in zip(r, align)] for i, r in enumerate(rows)]
            t = Table(data, repeatRows=1, colWidths=_col_widths(rows, width))
            style = [("GRID", (0, 0), (-1, -1), 0.5, line), ("BACKGROUND", (0, 0), (-1, 0), head_bg),
                     ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                     ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
            t.setStyle(TableStyle(style))
            story += [t, Spacer(1, 4)]
        elif kind == "bullet":
            story.append(Paragraph(rich(body), st["bullet"], bulletText="•"))
        else:
            story.append(Paragraph(rich(body), st[kind]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=parse(md)[0][1], author="영업 성과 분석 대시보드 (PoC)")
    doc.build(story)
    return buf.getvalue()


def _col_widths(rows, total):
    """열 내용 길이에 비례한 폭 (너무 좁거나 넓지 않게)."""
    def w(text):  # 한글은 영문·숫자의 약 2배 폭
        return sum(2 if ord(ch) > 0x2E7F else 1 for ch in BOLD.sub(r"\1", text))
    lens = [max(w(r[i]) for r in rows) for i in range(len(rows[0]))]
    lens = [min(max(n, 6), 50) + 2 for n in lens]
    return [total * n / sum(lens) for n in lens]


# ---------------------------------------------------------------- 워드 (python-docx)
def report_docx(md: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        s = doc.styles[name]
        s.font.name = "맑은 고딕"
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Heading 2"].font.color.rgb = RGBColor(0x15, 0x7F, 0x8D)

    def add_runs(p, text, size=None, color=None):
        for i, part in enumerate(BOLD.split(text)):  # 홀수 번째 조각이 **굵게**
            if part:
                r = p.add_run(part)
                r.bold = i % 2 == 1 or None
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = color

    for kind, body in parse(md):
        if kind == "h1":
            p = doc.add_heading(level=1)
            add_runs(p, body, 16)
        elif kind == "h2":
            p = doc.add_heading(level=2)
            add_runs(p, body)
        elif kind == "bullet":
            add_runs(doc.add_paragraph(style="List Bullet"), body)
        elif kind == "note":
            add_runs(doc.add_paragraph(), body, 9, RGBColor(0x5B, 0x6B, 0x78))
        elif kind == "para":
            add_runs(doc.add_paragraph(), body)
        else:
            rows, align = body["rows"], body["align"] or ["left"] * len(body["rows"][0])
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            widths = _col_widths(rows, Cm(17).emu)
            for j, col in enumerate(t.columns):
                col.width = int(widths[j])
            for i, r in enumerate(rows):
                for j, (c, a) in enumerate(zip(r, align)):
                    t.cell(i, j).width = int(widths[j])
                    p = t.cell(i, j).paragraphs[0]
                    add_runs(p, f"**{c}**" if i == 0 and c else c, 9,
                             RGBColor(0xD0, 0x3A, 0x3A) if "▼" in c or "⚠" in c else None)
                    if a == "right" and i:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
